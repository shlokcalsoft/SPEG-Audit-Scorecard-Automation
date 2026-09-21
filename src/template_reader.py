from pathlib import Path
import json
from openpyxl import load_workbook


# ---------------------------------------------------------
# PROJECT PATHS
# ---------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = (
    PROJECT_ROOT
    / "template"
    / "SEPG_Master_Audit_Template.xlsx"
)

REGISTRY_PATH = (
    PROJECT_ROOT
    / "data"
    / "checkpoint_registry.json"
)

SCORECARD_SHEET = "Q2 2026"


# ---------------------------------------------------------
# LOAD EXCEL TEMPLATE
# ---------------------------------------------------------

def load_template(template_path=TEMPLATE_PATH):

    print("\n[1/4] Checking template...", flush=True)

    if not template_path.exists():
        raise FileNotFoundError(
            f"Template not found: {template_path}"
        )

    print(f"      Found: {template_path}", flush=True)

    print("[2/4] Loading Excel workbook...", flush=True)

    workbook = load_workbook(
        template_path,
        data_only=False
    )

    print("      Workbook loaded successfully.", flush=True)

    return workbook


# ---------------------------------------------------------
# READ CHECKPOINTS
# ---------------------------------------------------------

def read_checkpoints(workbook):

    print("[3/4] Reading checkpoints...", flush=True)

    worksheet = workbook[SCORECARD_SHEET]

    checkpoints = []

    current_phase = None
    current_activity = None

    for row in range(1, worksheet.max_row + 1):

        # ---------------------------------------------
        # Read Phase
        # ---------------------------------------------

        phase_value = worksheet.cell(row, 2).value

        if phase_value is not None:

            phase_value = str(phase_value).strip()

            if phase_value:
                current_phase = phase_value


        # ---------------------------------------------
        # Read Activity
        # ---------------------------------------------

        activity_value = worksheet.cell(row, 3).value

        if activity_value is not None:

            activity_value = str(activity_value).strip()

            if activity_value:
                current_activity = activity_value


        # ---------------------------------------------
        # Read Checkpoint
        # ---------------------------------------------

        checkpoint_value = worksheet.cell(row, 4).value

        if checkpoint_value is None:
            continue

        checkpoint_text = str(checkpoint_value).strip()

        if not checkpoint_text:
            continue

        # Skip table header
        if checkpoint_text.lower() == "checkpoint":
            continue


        # ---------------------------------------------
        # Create checkpoint record
        # ---------------------------------------------

        checkpoint = {

            "checkpoint_id": f"C{len(checkpoints) + 1:03d}",

            "phase": current_phase,

            "activity": current_activity,

            "checkpoint_text": checkpoint_text,

            # Column M
            "weight": worksheet.cell(row, 13).value,

            # Excel mapping
            "sheet": SCORECARD_SHEET,

            "row": row,

            # Column L
            "response_cell": f"L{row}",

            # Column N
            "score_cell": f"N{row}",

            # Column O
            "weighted_score_cell": f"O{row}",

            # Column Q
            "observation_cell": f"Q{row}"
        }

        checkpoints.append(checkpoint)


    return checkpoints


# ---------------------------------------------------------
# SAVE CHECKPOINT REGISTRY
# ---------------------------------------------------------

def save_registry(checkpoints, registry_path=REGISTRY_PATH):

    print("[4/4] Creating checkpoint registry...", flush=True)

    # Make sure data folder exists
    registry_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    registry = {

        "template": {
            "name": TEMPLATE_PATH.name,
            "scorecard_sheet": SCORECARD_SHEET,
            "total_checkpoints": len(checkpoints)
        },

        "checkpoints": checkpoints
    }


    with open(
        registry_path,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            registry,
            file,
            indent=4,
            ensure_ascii=False
        )


    print(
        f"      Registry saved to:\n      {registry_path}",
        flush=True
    )


# ---------------------------------------------------------
# MAIN
# ---------------------------------------------------------

def main():

    print("=" * 60)
    print("SEPG TEMPLATE READER")
    print("=" * 60)

    print("\nTemplate path:")
    print(TEMPLATE_PATH)


    # Load workbook
    workbook = load_template()


    # Show sheets
    print("\nSheets found:")

    for sheet in workbook.sheetnames:
        print(f"  - {sheet}")


    # Read checkpoints
    checkpoints = read_checkpoints(workbook)


    print(
        f"\nTotal checkpoints found: {len(checkpoints)}"
    )


    # Validate expected count
    if len(checkpoints) != 248:

        print(
            "\nWARNING: Expected 248 checkpoints "
            f"but found {len(checkpoints)}."
        )

    else:

        print(
            "\nCheckpoint count validation: PASSED"
        )


    # Display first 5 checkpoints
    print("\nFirst 5 checkpoints:")
    print("-" * 60)

    for checkpoint in checkpoints[:5]:
        print(checkpoint)


    # Save registry
    save_registry(checkpoints)


    print("\n" + "=" * 60)
    print("Template reading completed successfully.")
    print("=" * 60)


if __name__ == "__main__":
    main()