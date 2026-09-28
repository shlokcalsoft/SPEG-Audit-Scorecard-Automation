import json

from openpyxl import Workbook, load_workbook

from src.excel_exporter import export_assessments


def test_export_fills_response_and_observation_preserving_score_formulas(tmp_path):
    template_path = tmp_path / "template.xlsx"
    registry_path = tmp_path / "registry.json"
    assessments_path = tmp_path / "assessments.json"
    output_path = tmp_path / "completed.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Q2 2026"
    worksheet["N12"] = '=IF(L12="Yes",5,0)'
    worksheet["O12"] = "=N12*M12"
    workbook.save(template_path)

    registry_path.write_text(json.dumps({"checkpoints": [{
        "checkpoint_id": "C001", "sheet": "Q2 2026", "response_cell": "L12",
        "score_cell": "N12", "weighted_score_cell": "O12", "observation_cell": "Q12",
    }]}), encoding="utf-8")
    assessments_path.write_text(json.dumps({"assessments": [{
        "checkpoint_id": "C001", "response": "Yes", "observation": "The owner confirmed it.",
        "needs_review": True, "review_reason": "Confirm the linked artifact.",
    }]}), encoding="utf-8")

    export_assessments(template_path, registry_path, assessments_path, output_path)
    exported = load_workbook(output_path, data_only=False)["Q2 2026"]

    assert exported["L12"].value == "Yes"
    assert exported["N12"].value == '=IF(L12="Yes",5,0)'
    assert exported["O12"].value == "=N12*M12"
    assert exported["Q12"].value == "[C001] The owner confirmed it.\nREVIEW REQUIRED: Confirm the linked artifact."
    assert exported["L12"].fill.fgColor.rgb.endswith("FFF2CC")


def test_export_writes_observation_to_merged_range_anchor(tmp_path):
    template_path = tmp_path / "merged_template.xlsx"
    registry_path = tmp_path / "merged_registry.json"
    assessments_path = tmp_path / "merged_assessments.json"
    output_path = tmp_path / "merged_completed.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Q2 2026"
    worksheet.merge_cells("Q12:S13")
    workbook.save(template_path)

    registry_path.write_text(json.dumps({"checkpoints": [{
        "checkpoint_id": "C001", "sheet": "Q2 2026", "response_cell": "L12",
        "score_cell": "N12", "weighted_score_cell": "O12", "observation_cell": "Q13",
    }]}), encoding="utf-8")
    assessments_path.write_text(json.dumps({"assessments": [{
        "checkpoint_id": "C001", "response": "Partial", "observation": "Some information was discussed.",
        "needs_review": False,
    }]}), encoding="utf-8")

    export_assessments(template_path, registry_path, assessments_path, output_path)
    exported = load_workbook(output_path, data_only=False)["Q2 2026"]

    assert exported["Q12"].value == "[C001] Some information was discussed."


def test_export_hides_raw_provider_errors_from_observation(tmp_path):
    template_path = tmp_path / "provider_template.xlsx"
    registry_path = tmp_path / "provider_registry.json"
    assessments_path = tmp_path / "provider_assessments.json"
    output_path = tmp_path / "provider_completed.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Q2 2026"
    workbook.save(template_path)
    registry_path.write_text(json.dumps({"checkpoints": [{
        "checkpoint_id": "C001", "sheet": "Q2 2026", "response_cell": "L12",
        "score_cell": "N12", "weighted_score_cell": "O12", "observation_cell": "Q12",
    }]}), encoding="utf-8")
    assessments_path.write_text(json.dumps({"assessments": [{
        "checkpoint_id": "C001", "response": "Not Clarified", "observation": "No validated quote.",
        "needs_review": True,
        "review_reason": "Provider request failed: All configured OpenRouter models failed: HTTP 402 secret detail",
    }]}), encoding="utf-8")

    export_assessments(template_path, registry_path, assessments_path, output_path)
    exported = load_workbook(output_path, data_only=False)["Q2 2026"]

    assert exported["Q12"].value == "[C001] No validated quote.\nREVIEW REQUIRED: AI service unavailable; review this checkpoint manually."
    assert "HTTP 402" not in exported["Q12"].value